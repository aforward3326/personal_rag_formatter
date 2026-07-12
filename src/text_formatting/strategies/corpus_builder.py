import os
import re
import hashlib
import platform
from pathlib import Path
from datetime import datetime, timezone

from .base import TextFormatterStrategy
from ..utils.file_utils import save_json_batch, load_progress, save_progress, clear_progress, prompt_resume

# Try to import parsing libraries and set flags
try:
    import docx
    HAS_DOCX = True
except ImportError: HAS_DOCX = False
try:
    import pptx
    HAS_PPTX = True
except ImportError: HAS_PPTX = False
try:
    import pdfplumber
    import fitz  # PyMuPDF
    HAS_PDF = True
except ImportError: HAS_PDF = False
try:
    import easyocr
    HAS_OCR = True
except ImportError: HAS_OCR = False

class CorpusBuilder(TextFormatterStrategy):
    """
    Processes a directory of various document types (.txt, .docx, .pptx, .pdf),
    extracts text content (using OCR for image-based PDFs), and builds a
    unified corpus for RAG ingestion.
    """

    def __init__(self, input_dir: Path, output_file: Path, **kwargs):
        super().__init__(input_dir, output_file, **kwargs)
        self.ocr_reader = None
        self._check_dependencies()

    def _check_dependencies(self):
        if not HAS_DOCX: self.logger.warning("python-docx not found. .docx files will be skipped.")
        if not HAS_PPTX: self.logger.warning("python-pptx not found. .pptx files will be skipped.")
        if not HAS_PDF: self.logger.warning("pdfplumber or PyMuPDF not found. .pdf files will be skipped.")
        if not HAS_OCR: self.logger.warning("easyocr not found. OCR for PDFs will be disabled.")

    def _init_ocr_reader(self):
        if self.ocr_reader is None and HAS_OCR:
            self.logger.info("Initializing EasyOCR (ch_tra + en)...")
            self.ocr_reader = easyocr.Reader(['ch_tra', 'en'], gpu=False)

    def process(self):
        self.logger.info("Process Started: Corpus Builder")
        if not self.input_dir.exists():
            self.logger.error(f"Source data directory not found: {self.input_dir}")
            return

        progress_data = load_progress(self.progress_file)
        if not prompt_resume(progress_data, self.logger):
            self.logger.info("Starting fresh, clearing previous progress.")
            clear_progress(self.progress_file)
            progress_data = {"processed_files": set()}
        else:
            self.logger.info("Resuming process...")

        processed_files = set(progress_data.get("processed_files", []))
        batch_results = []
        
        try:
            all_files = list(self.input_dir.rglob('*'))
            self.logger.info(f"Found {len(all_files)} files to scan.")

            for file_path in all_files:
                if file_path.is_dir(): continue
                
                rel_path = str(file_path.relative_to(self.input_dir))
                if rel_path in processed_files:
                    continue

                ext = file_path.suffix.lower().lstrip('.')
                parser = self._get_parser(ext)
                if not parser:
                    continue

                self.logger.info(f"Processing: {rel_path}")
                try:
                    text, is_ocr = parser(file_path)
                    cleaned_text = self._clean_text(text)
                    if cleaned_text:
                        batch_results.append(self._create_corpus_entry(file_path, rel_path, cleaned_text, is_ocr))
                except Exception as e:
                    self.logger.error(f"Error processing file {rel_path}: {e}")

                processed_files.add(rel_path)
                if len(batch_results) >= 100: # Using a fixed batch size
                    save_json_batch(batch_results, self.output_dir, self.output_file.name, "corpus")
                    progress_data["processed_files"] = list(processed_files)
                    save_progress(self.progress_file, progress_data)
                    batch_results = []

            if batch_results:
                save_json_batch(batch_results, self.output_dir, self.output_file.name, "corpus")

            self.logger.info("Process Finished Successfully.")
            clear_progress(self.progress_file)

        except KeyboardInterrupt:
            progress_data["processed_files"] = list(processed_files)
            self.logger.warning("Process interrupted. Progress saved.")
            save_progress(self.progress_file, progress_data)
        except Exception as e:
            progress_data["processed_files"] = list(processed_files)
            self.logger.error(f"An unexpected error occurred: {e}", exc_info=True)
            save_progress(self.progress_file, progress_data)

    def _get_parser(self, extension: str):
        if extension == 'txt': return self._parse_txt
        if extension == 'docx' and HAS_DOCX: return self._parse_docx
        if extension == 'pptx' and HAS_PPTX: return self._parse_pptx
        if extension == 'pdf' and HAS_PDF: return self._parse_pdf
        return None

    def _clean_text(self, text: str) -> str:
        if not text: return ""
        text = re.sub(r'\n+', '\n', text)
        return re.sub(r'[ \t]+', ' ', text).strip()

    def _create_corpus_entry(self, path: Path, rel_path: str, text: str, is_ocr: bool) -> dict:
        return {
            "doc_id": hashlib.md5(rel_path.encode('utf-8')).hexdigest(),
            "file_name": path.name,
            "page_content": text,
            "metadata": {
                "doc_type": self._get_doc_type(rel_path),
                "file_format": path.suffix.lower().lstrip('.'),
                "created_at": self._get_creation_time(path),
                "is_ocr": is_ocr,
                "relative_path": rel_path
            }
        }

    def _get_creation_time(self, path: Path) -> str:
        try:
            stat = path.stat()
            ts = stat.st_birthtime if hasattr(stat, 'st_birthtime') else stat.st_mtime
            return datetime.fromtimestamp(ts, tz=timezone.utc).isoformat()
        except: return ""

    def _get_doc_type(self, path_str: str) -> str:
        lower_path = path_str.lower()
        if any(kw in lower_path for kw in ['學術', '論文', 'paper']): return 'academic_paper'
        if any(kw in lower_path for kw in ['日記', '日誌', 'diary']): return 'journal_diary'
        if any(kw in lower_path for kw in ['報告', 'report', 'meeting']): return 'business_report'
        if any(kw in lower_path for kw in ['寫作', '小說', 'novel']): return 'creative_writing'
        return 'general'

    def _parse_txt(self, path: Path) -> tuple:
        with open(path, 'r', encoding='utf-8', errors='replace') as f:
            return f.read(), False

    def _parse_docx(self, path: Path) -> tuple:
        doc = docx.Document(path)
        text = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip(): text.append(cell.text)
        return "\n".join(text), False

    def _parse_pptx(self, path: Path) -> tuple:
        prs = pptx.Presentation(path)
        text = [shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        return "\n".join(text), False



    def _parse_pdf(self, path: Path) -> tuple:
        text = ""
        try:
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text: text += page_text + "\n"
        except Exception as e:
            self.logger.warning(f"pdfplumber direct parsing failed for {path.name}: {e}")

        if len(text.strip()) < 50:
            self.logger.info(f"Low text content, triggering OCR for {path.name}.")
            self._init_ocr_reader()
            if not self.ocr_reader:
                self.logger.error("OCR reader not available, cannot process image-based PDF.")
                return "", False
            
            ocr_text = ""
            try:
                doc = fitz.open(path)
                for page in doc:
                    pix = page.get_pixmap(dpi=150)
                    results = self.ocr_reader.readtext(pix.tobytes("png"), detail=0)
                    ocr_text += " ".join(results) + "\n"
                doc.close()
                return ocr_text, True
            except Exception as e:
                self.logger.error(f"OCR parsing failed for {path.name}: {e}")
                return "", False
        
        return text, False
