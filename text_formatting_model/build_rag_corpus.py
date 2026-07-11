import os
import json
import hashlib
import datetime
import logging
import re
from pathlib import Path
import platform
import sys
from dotenv import load_dotenv

# Load environment variables from .env file
load_dotenv('../pipeline.env')

# --- Configuration from environment variables ---
BASE_DIR = os.getenv("CORPUS_INPUT_DIR", "../org_data/word_data")
OUTPUT_FILE = os.getenv("CORPUS_OUTPUT_FILE", "../output_clean_data/personal_corpus_rag_ready.json")
LOG_DIR = os.getenv("LOG_DIR", "../log")
PROC_DIR = os.getenv("PROC_DIR", "../processing")
PROCESS_NAME = "build_rag_corpus"

# Ensure directories exist
Path(LOG_DIR).mkdir(parents=True, exist_ok=True)
Path(PROC_DIR).mkdir(parents=True, exist_ok=True)
process_log_dir = Path(LOG_DIR) / PROCESS_NAME
process_log_dir.mkdir(parents=True, exist_ok=True)
process_proc_dir = Path(PROC_DIR) / PROCESS_NAME
process_proc_dir.mkdir(parents=True, exist_ok=True)

# Setup logging
def setup_logger():
    date_str = datetime.datetime.now().strftime("%Y%m%d")
    base_log_filename = f"{PROCESS_NAME}_{date_str}.log"
    log_file_path = process_log_dir / base_log_filename

    # Handle max 1GB log size and batching
    batch = 1
    while log_file_path.exists() and log_file_path.stat().st_size > 1 * 1024 * 1024 * 1024:
        log_file_path = process_log_dir / f"{PROCESS_NAME}_{date_str}_{batch}.log"
        batch += 1

    logger = logging.getLogger(PROCESS_NAME)
    logger.setLevel(logging.INFO)
    
    if not logger.handlers:
        file_handler = logging.FileHandler(log_file_path, encoding='utf-8')
        formatter = logging.Formatter('%(asctime)s - %(levelname)s - [%(processName)s] - %(message)s')
        file_handler.setFormatter(formatter)
        logger.addHandler(file_handler)
        
        console_handler = logging.StreamHandler()
        console_handler.setFormatter(formatter)
        logger.addHandler(console_handler)

    return logger

logger = setup_logger()

# Setup Progress tracking
PROGRESS_FILE = process_proc_dir / "progress.json"

def load_progress():
    if PROGRESS_FILE.exists():
        try:
            with open(PROGRESS_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception as e:
            logger.error(f"Error loading progress: {e}")
    return {"processed_files": [], "total_files": 0}

def save_progress(progress_data):
    try:
        with open(PROGRESS_FILE, 'w', encoding='utf-8') as f:
            json.dump(progress_data, f, ensure_ascii=False, indent=2)
    except Exception as e:
        logger.error(f"Error saving progress: {e}")

def clear_progress():
    if PROGRESS_FILE.exists():
        PROGRESS_FILE.unlink()

def prompt_resume(progress_data):
    if progress_data.get("processed_files"):
        while True:
            response = input(f"Found interrupted progress ({len(progress_data['processed_files'])} files processed). Do you want to resume? (y/n): ").strip().lower()
            if response in ['y', 'n']:
                return response == 'y'
    return False

# --- Check required modules ---
try:
    import docx
    import pptx
    import pdfplumber
    import fitz  # PyMuPDF
    import easyocr
except ImportError as e:
    logger.error(f"Missing required package: {e}. Please run pip install -r requirements.txt")

# Lazy initialize EasyOCR
ocr_reader = None

def get_output_file_path(output_path: str) -> Path:
    """Generate output file path, appending timestamp and batch number if exists"""
    base_path = Path(output_path)
    if not base_path.exists():
        return base_path
        
    timestamp = datetime.datetime.now().strftime("%Y%m%d%H")
    name = base_path.stem
    ext = base_path.suffix
    directory = base_path.parent
    
    base_new_filename = f"{name}_{timestamp}{ext}"
    output_new_path = directory / base_new_filename
    
    if not output_new_path.exists():
        return output_new_path
    
    batch = 1
    while True:
        new_filename = f"{name}_{timestamp}_{batch}{ext}"
        new_path = directory / new_filename
        if not new_path.exists():
            return new_path
        batch += 1

def save_batch(final_data, output_path_dir, base_filename):
    output_file_path = get_output_file_path(output_path_dir / base_filename)
    try:
        with open(output_file_path, 'w', encoding='utf-8') as f:
            json.dump({"corpus": final_data}, f, ensure_ascii=False, indent=2)
        logger.info(f"Successfully saved batch to {output_file_path}")
    except Exception as e:
        logger.error(f"Error saving batch JSON file: {e}")

def get_ocr_reader():
    global ocr_reader
    if ocr_reader is None:
        logger.info("First OCR trigger, initializing EasyOCR (ch_tra + en)...")
        # gpu=False if no dedicated GPU, change to True for better performance if applicable
        ocr_reader = easyocr.Reader(['ch_tra', 'en'], gpu=False)
    return ocr_reader

# --- Text processing and utilities ---
def clean_text(text):
    """Clean excessive newlines and spaces"""
    if not text:
        return ""
    text = re.sub(r'\n+', '\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    return text.strip()

def get_md5(string):
    """Calculate MD5 hash of a string as UUID"""
    return hashlib.md5(string.encode('utf-8')).hexdigest()

def format_time(timestamp):
    """Convert OS Timestamp to ISO-8601 (UTC)"""
    dt = datetime.datetime.fromtimestamp(timestamp, datetime.timezone.utc)
    return dt.strftime('%Y-%m-%dT%H:%M:%SZ')

def get_creation_time(file_path_str):
    """Get file creation time compatible with different OS"""
    if platform.system() == 'Windows':
        return os.path.getctime(file_path_str)
    else:
        stat = os.stat(file_path_str)
        try:
            return stat.st_birthtime
        except AttributeError:
            # Fallback to mtime on systems that do not support st_birthtime
            return stat.st_mtime

def get_doc_type(file_path_str):
    """Auto-categorize document type based on path and filename heuristics"""
    lower_path = file_path_str.lower()
    if any(kw in lower_path for kw in ['學術', '論文', 'paper', 'academic', 'research', 'thesis', 'dissertation']):
        return 'academic_paper'
    if any(kw in lower_path for kw in ['日記', '日誌', 'diary', 'journal']):
        return 'journal_diary'
    if any(kw in lower_path for kw in ['出差', '報告', 'business', 'report', 'meeting', '會議']):
        return 'business_report'
    if any(kw in lower_path for kw in ['寫作', '小說', 'creative', 'writing', 'novel', 'story', '故事']):
        return 'creative_writing'
    return 'general'

# --- Parsing functions ---
def parse_txt(file_path):
    with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
        return f.read(), False

def parse_docx(file_path):
    doc = docx.Document(file_path)
    text_blocks = []
    # Extract paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            text_blocks.append(para.text)
    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_blocks.append(cell.text)
    return "\n".join(text_blocks), False

def parse_pptx(file_path):
    prs = pptx.Presentation(file_path)
    text_blocks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_blocks.append(shape.text)
    return "\n".join(text_blocks), False

def parse_pdf(file_path):
    text = ""
    is_ocr = False
    
    # 1. Try extracting text directly (pdfplumber)
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        logger.warning(f"pdfplumber direct parsing failed for {file_path}: {e}")

    # 2. OCR trigger condition: extracted text length is less than 50 chars
    if len(text.strip()) < 50:
        logger.info(f"Detected as scanned or image PDF (text length {len(text.strip())}), triggering OCR for {file_path}.")
        is_ocr = True
        text = ""
        try:
            reader = get_ocr_reader()
            doc = fitz.open(file_path)
            for page in doc:
                # Convert to image (resolution dpi=150)
                pix = page.get_pixmap(dpi=150)
                img_bytes = pix.tobytes("png")
                # Parse image byte with easyocr
                results = reader.readtext(img_bytes, detail=0)
                text += " ".join(results) + "\n"
            doc.close()
        except Exception as e:
            logger.error(f"OCR parsing failed for {file_path}: {e}")

    return text, is_ocr

# --- Main Logic ---
def main():
    start_time = datetime.datetime.now()
    logger.info("Process Started")
    base_path = Path(BASE_DIR)
    
    output_path = Path(OUTPUT_FILE)
    output_dir = output_path.parent
    output_dir.mkdir(parents=True, exist_ok=True)
    base_filename = output_path.name
    
    if not base_path.exists():
        logger.error(f"Source data directory not found: {BASE_DIR}. Please check the path.")
        return

    progress_data = load_progress()
    resume = prompt_resume(progress_data)
    
    if not resume:
        logger.info("Starting fresh, clearing previous progress.")
        clear_progress()
        progress_data = {"processed_files": [], "total_files": 0}
    else:
        logger.info(f"Resuming process. {len(progress_data.get('processed_files', []))} files already processed.")

    processed_files = progress_data.get("processed_files", [])
    total_files = progress_data.get("total_files", 0)

    logger.info(f"Start scanning and processing directory: {BASE_DIR}")
    
    batch_results = []
    
    try:
        for root, _, files in os.walk(base_path):
            for file in files:
                file_path = Path(root) / file
                ext = file_path.suffix.lower().lstrip('.')
                
                # Filter unsupported formats
                if ext not in ['txt', 'docx', 'pptx', 'pdf']:
                    continue

                rel_path = file_path.relative_to(base_path).as_posix()
                if rel_path in processed_files:
                    continue

                total_files += 1
                doc_id = get_md5(rel_path)
                creation_time = format_time(get_creation_time(file_path))
                doc_type = get_doc_type(rel_path)

                logger.info(f"Processing: {rel_path} ({ext})")
                
                text = ""
                is_ocr = False
                
                try:
                    if ext == 'txt':
                        text, is_ocr = parse_txt(file_path)
                    elif ext == 'docx':
                        text, is_ocr = parse_docx(file_path)
                    elif ext == 'pptx':
                        text, is_ocr = parse_pptx(file_path)
                    elif ext == 'pdf':
                        text, is_ocr = parse_pdf(file_path)
                    
                    # Clean and skip empty
                    cleaned_text = clean_text(text)
                    if not cleaned_text:
                        logger.warning(f"File is empty or unable to extract text, skipping: {rel_path}")
                    else:
                        entry = {
                            "doc_id": doc_id,
                            "file_name": file_path.name,
                            "page_content": cleaned_text,
                            "metadata": {
                                "doc_type": doc_type,
                                "file_format": ext,
                                "created_at": creation_time,
                                "is_ocr": is_ocr,
                                "relative_path": rel_path
                            }
                        }
                        batch_results.append(entry)

                except Exception as e:
                    logger.error(f"Unexpected error processing file {rel_path}: {e}")

                processed_files.append(rel_path)
                progress_data["processed_files"] = processed_files
                progress_data["total_files"] = total_files
                save_progress(progress_data)

                if len(batch_results) >= 100:
                    save_batch(batch_results, output_dir, base_filename)
                    batch_results = []

        if batch_results:
            save_batch(batch_results, output_dir, base_filename)

        end_time = datetime.datetime.now()
        execution_time = end_time - start_time

        logger.info("-" * 40)
        logger.info("Corpus building completed - Final Report:")
        logger.info(f"Total files processed: {total_files}")
        logger.info(f"Total Execution Time: {execution_time}")
        logger.info("-" * 40)

        clear_progress()
        logger.info("Process Finished Successfully.")

    except KeyboardInterrupt:
        logger.warning("Process interrupted by user.")
        save_progress(progress_data)
        sys.exit(0)
    except Exception as e:
        logger.error(f"Unexpected error occurred: {e}", exc_info=True)
        save_progress(progress_data)
        sys.exit(1)

if __name__ == "__main__":
    main()